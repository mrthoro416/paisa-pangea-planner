"""Offline import of the supplied matrix; charts are deliberately ignored.
Usage: python3 scripts/import-visa-matrix.py path/to/workbook.xlsx
Requires openpyxl. Output is auditable source data, never a browser XLSX parser.
"""
import hashlib
import json
import sys
from pathlib import Path
import openpyxl

source = Path(sys.argv[1])
workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
passports = {}
destinations = {}
records = {}
missing_template = None
missing_count = 0
for sheet_name in workbook.sheetnames[2:]:
    destination_code = sheet_name.rsplit("_", 1)[1]
    header = None
    codes = set()
    for row in workbook[sheet_name].iter_rows(values_only=True):
        if row and row[0] == "passport_code":
            header = list(row[:16])
            continue
        if not header or not row or row[0] is None:
            continue
        record = dict(zip(header, row[:16]))
        code = record["passport_code"]
        if code in codes:
            raise ValueError(f"Duplicate pair {code}:{destination_code}")
        codes.add(code)
        passports[code] = {"code": code, "name": record["passport_country"], "iso3": record["passport_iso3"]}
        destinations[destination_code] = {"name": record["destination_authority_name"], "url": record["destination_authority_url"]}
        if record["coverage_status"] == "missing":
            missing_count += 1
            template = {k: v for k, v in record.items() if k not in ["passport_code", "passport_country", "passport_iso3", "destination_authority_name", "destination_authority_url"]}
            if missing_template is not None and template != missing_template:
                raise ValueError("Missing rows differ; expand the import format instead of discarding fields")
            missing_template = template
        else:
            record["destination_code"] = destination_code
            records[f"{code}:{destination_code}"] = record
    if len(codes) != 198:
        raise ValueError(f"Unexpected passport row count in {sheet_name}: {len(codes)}")
if len(destinations) != 198 or len(passports) != 198:
    raise ValueError("Workbook coverage changed; review import assumptions")
data = {"workbook": source.name, "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "format_version": 1, "total_pairs": len(destinations) * len(passports),
        "missing_count": missing_count, "passports": list(passports.values()),
        "destination_authorities": destinations, "missing_template": missing_template, "records": records}
Path("src/data/visa-matrix.json").write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n")
print(f"Imported {len(records)} explicit records and {missing_count} losslessly represented missing rows")
